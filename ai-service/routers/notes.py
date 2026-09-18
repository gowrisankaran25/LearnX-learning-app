from fastapi import APIRouter, HTTPException, UploadFile, File
from pydantic import BaseModel
from llm.client import llm_client
import os
import shutil
from typing import Optional

router = APIRouter()

# Create uploads directory if it doesn't exist
UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

async def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from PDF file"""
    try:
        import pypdf
        text = ""
        with open(file_path, 'rb') as file:
            pdf_reader = pypdf.PdfReader(file)
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
        return text
    except Exception as e:
        print(f"PDF extraction error: {e}")
        return ""

async def extract_text_from_pptx(file_path: str) -> str:
    """Extract text from PowerPoint file"""
    try:
        from pptx import Presentation
        text = ""
        prs = Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
        return text
    except Exception as e:
        print(f"PPTX extraction error: {e}")
        return ""

async def extract_text_from_txt(file_path: str) -> str:
    """Extract text from text file"""
    try:
        with open(file_path, 'r', encoding='utf-8') as file:
            return file.read()
    except Exception as e:
        print(f"TXT extraction error: {e}")
        return ""

@router.post("/")
async def analyze_notes(file: UploadFile = File(...)):
    """Analyze uploaded notes/PDF"""
    try:
        # Save uploaded file
        file_path = os.path.join(UPLOAD_DIR, file.filename)
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
        
        # Extract text based on file type
        text = ""
        if file.filename.endswith('.pdf'):
            text = await extract_text_from_pdf(file_path)
        elif file.filename.endswith('.pptx') or file.filename.endswith('.ppt'):
            text = await extract_text_from_pptx(file_path)
        elif file.filename.endswith('.txt'):
            text = await extract_text_from_txt(file_path)
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type")
        
        if not text:
            raise HTTPException(status_code=400, detail="Could not extract text from file")
        
        # Generate summary
        summary = await llm_client.summarize(text)
        
        # Generate key points (could be enhanced with LLM)
        key_points = text.split('\n')[:10]  # Simple extraction
        
        # Generate flashcards (simplified - would use LLM in production)
        flashcards = []
        
        # Generate questions (simplified - would use LLM in production)
        questions = []
        
        # Clean up uploaded file
        os.remove(file_path)
        
        return {
            "summary": summary,
            "keyPoints": key_points,
            "flashcards": flashcards,
            "questions": questions
        }
    except Exception as e:
        print(f"Notes analysis error: {e}")
        raise HTTPException(status_code=500, detail=str(e))
